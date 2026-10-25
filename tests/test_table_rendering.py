from __future__ import annotations

import csv
import io
from contextlib import nullcontext
from html.parser import HTMLParser

from streamlit.testing.v1 import AppTest

import app as evaluator
from candidate_evaluator.roles import get_role_profile


class _TableParser(HTMLParser):
    def __init__(self, body: str):
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self.tags: list[str] = []
        self.outside_text: list[str] = []
        self._cell: list[str] | None = None
        self._in_style = False
        self.feed(body)

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        if tag == "style":
            self._in_style = True
        elif tag == "tr":
            self.rows.append([])
        elif tag in {"td", "th"}:
            self._cell = []

    def handle_endtag(self, tag):
        if tag == "style":
            self._in_style = False
        elif tag in {"td", "th"} and self._cell is not None:
            self.rows[-1].append("".join(self._cell))
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)
        elif not self._in_style and data.strip():
            self.outside_text.append(data)


def _render(rows: list[dict]) -> AppTest:
    app = AppTest.from_string(
        "from app import _display_table\n"
        f"_display_table({rows!r})\n"
    ).run()
    assert not app.exception
    return app


def _table(app: AppTest) -> _TableParser:
    # The HTML delta bypasses Markdown entirely, so profile blank lines cannot
    # terminate its table or turn following markup into a Markdown code block.
    assert not app.markdown
    assert len(app.get("html")) == 1
    element = app.get("html")[0]
    assert not getattr(element.proto, "unsafe_allow_javascript", False)
    return _TableParser(element.proto.body)


def test_multiline_rationales_remain_in_each_candidates_table_cell() -> None:
    first = "First capability has direct evidence.\n\nSecond capability is unverified.\nOverall: supported by the fixed scores."
    second = "Separate candidate explanation.\n\nIts evidence remains separate."
    app = _render([
        {"Candidate": "First Candidate", "Final Score": 37.2, "Score Rationale": first},
        {"Candidate": "Second Candidate", "Final Score": 18.6, "Score Rationale": second},
    ])

    parsed = _table(app)

    assert parsed.rows == [
        ["Candidate", "Final Score", "Score Rationale"],
        ["First Candidate", "37.2", first],
        ["Second Candidate", "18.6", second],
    ]
    assert not parsed.outside_text


def test_hostile_html_and_markdown_are_displayed_as_literal_cell_text() -> None:
    hostile_header = '<svg onload="alert(1)">Candidate</svg>'
    hostile_value = '</td></tr><script>alert("injected")</script><img src=x onerror="alert(1)">'
    rationale = "Normal evidence.\n\n<table><tr><td>Injected row</td></tr></table>\n\n# Heading\n[Open](javascript:alert(1))"
    app = _render([
        {hostile_header: hostile_value, "Score Rationale": rationale},
        {hostile_header: "Safe Candidate", "Score Rationale": "Separate rationale."},
    ])

    parsed = _table(app)

    assert parsed.rows == [
        [hostile_header, "Score Rationale"],
        [hostile_value, rationale],
        ["Safe Candidate", "Separate rationale."],
    ]
    assert parsed.tags.count("table") == 1
    assert not {"script", "img", "svg", "a", "h1"}.intersection(parsed.tags)
    assert not parsed.outside_text


def test_sparse_rows_keep_values_aligned_with_their_column() -> None:
    parsed = _table(_render([
        {"Candidate": "First", "Score Rationale": "First explanation."},
        {"Candidate": "Second"},
    ]))

    assert parsed.rows[2] == ["Second", ""]


def test_empty_rows_show_a_caption_without_table_markup() -> None:
    app = _render([])

    assert app.caption[0].value == "No rows to show."
    assert not app.get("html")
    assert not app.markdown


def test_csv_excel_and_preview_use_current_prepared_rows_instead_of_a_saved_workbook(tmp_path, monkeypatch) -> None:
    from openpyxl import load_workbook

    monkeypatch.chdir(tmp_path)
    run_id = "saved-run"
    role = get_role_profile("sde3")
    stale = tmp_path / "outputs" / f"candidate_evaluation_sde3_{run_id}.xlsx"
    stale.parent.mkdir()
    stale.write_bytes(b"Stale workbook with an obsolete explanation.")
    source = [{"Candidate": "Current Candidate", "Score Rationale": "Old unbounded source prose."}]
    fixed = [{"Candidate": "Current Candidate", "Score Rationale": "Current normalized explanation."}]
    downloads = {}
    previews = []

    def capture_download(label, data, **kwargs):
        downloads[label] = data

    monkeypatch.setattr(evaluator, "role_for_run", lambda _: role)
    monkeypatch.setattr(evaluator, "result_rows", lambda _: source)
    monkeypatch.setattr(evaluator, "progress_counts", lambda _: {})
    monkeypatch.setattr(evaluator, "prepare_output_rows", lambda rows, selected_role: fixed)
    monkeypatch.setattr(evaluator, "_display_table", previews.append)
    monkeypatch.setattr(evaluator.st, "subheader", lambda *_: None)
    monkeypatch.setattr(evaluator.st, "columns", lambda _: [nullcontext(), nullcontext()])
    monkeypatch.setattr(evaluator.st, "download_button", capture_download)

    evaluator._show_downloads(run_id)

    csv_rows = list(csv.DictReader(io.StringIO(downloads["Download CSV"].decode("utf-8"))))
    workbook = load_workbook(io.BytesIO(downloads["Download Excel"]))
    excel_rows = list(workbook.active.iter_rows(values_only=True))
    excel_row = dict(zip(excel_rows[0], excel_rows[1]))
    assert csv_rows[0]["Score Rationale"] == "Current normalized explanation."
    assert excel_row["Score Rationale"] == csv_rows[0]["Score Rationale"]
    assert excel_row["Candidate"] == "Current Candidate"
    assert previews == [fixed]
    assert stale.read_bytes() == b"Stale workbook with an obsolete explanation."


def test_mismatched_rubric_disables_new_sde3_evaluation() -> None:
    app = AppTest.from_string(
        "import app as evaluator\n"
        "from unittest.mock import patch\n"
        "with patch.object(evaluator, 'run_select_options', return_value={}):\n"
        "    evaluator._show_run_controls([{'candidate_name': 'Candidate'}], "
        "'Technical assessment with maximum score 100', 'sde3', None, [], "
        "'openai', 'mock-model', 'none', 1, 'mock-key', True)\n"
    ).run()

    assert not app.exception
    assert app.button[0].label == "Start / resume evaluation"
    assert app.button[0].disabled
    assert "does not match SDE3 LinkedIn scoring" in app.error[0].value


def test_saved_incompatible_rubric_stops_before_resuming_or_modifying_progress(tmp_path) -> None:
    (tmp_path / "rubric.md").write_text("Technical assessment with maximum score 100", encoding="utf-8")
    app = AppTest.from_string(
        "from pathlib import Path\n"
        "from unittest.mock import patch\n"
        "import app as evaluator\n"
        "from candidate_evaluator.roles import get_role_profile\n"
        "with patch.object(evaluator, 'load_candidates', return_value=[]), "
        "patch.object(evaluator, 'role_for_run', return_value=get_role_profile('sde3')), "
        f"patch.object(evaluator, 'run_dir', return_value=Path({str(tmp_path)!r})), "
        "patch.object(evaluator, 'recover_batch_results', side_effect=AssertionError('Must not modify progress')):\n"
        "    evaluator._run_evaluation('saved-run', 'mock-key', 1, False)\n"
    ).run()

    assert not app.exception
    assert "does not match SDE3 LinkedIn scoring" in app.error[0].value


def test_saved_incompatible_rubric_disables_resume_and_retry_but_keeps_exports(tmp_path) -> None:
    (tmp_path / "rubric.md").write_text("Technical assessment with maximum score 100", encoding="utf-8")
    app = AppTest.from_string(
        "from pathlib import Path\n"
        "from unittest.mock import patch\n"
        "import app as evaluator\n"
        "from candidate_evaluator.roles import get_role_profile\n"
        "rubric = Path(evaluator._default_rubric_path('sde3')).read_text(encoding='utf-8')\n"
        "with patch.object(evaluator, 'run_select_options', return_value={'Saved run': 'saved-run'}), "
        "patch.object(evaluator, 'load_status', return_value={'role': 'sde3', 'model': 'mock-model'}), "
        "patch.object(evaluator, 'load_candidates', return_value=[]), "
        "patch.object(evaluator, 'role_for_run', return_value=get_role_profile('sde3')), "
        f"patch.object(evaluator, 'run_dir', return_value=Path({str(tmp_path)!r})), "
        "patch.object(evaluator, '_candidate_input_issues_for_role', return_value=[]), "
        "patch.object(evaluator, '_show_evaluation_status'), patch.object(evaluator, '_show_downloads'):\n"
        "    evaluator._show_run_controls([{'candidate_name': 'Candidate'}], rubric, 'sde3', None, [], "
        "'openai', 'mock-model', 'none', 1, 'mock-key', True)\n"
    ).run()
    app.selectbox[0].set_value("Saved run").run()

    assert not app.exception
    assert app.button[0].disabled
    assert app.button[1].disabled
    assert not app.button[2].disabled
    assert "does not match SDE3 LinkedIn scoring" in app.error[0].value
